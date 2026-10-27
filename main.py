import pandas as pd
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, FileResponse
from pydantic import BaseModel
import os
from datetime import datetime
import uvicorn
from typing import List
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment
import smtplib
from email.message import EmailMessage
import mimetypes

app = FastAPI(title="API de Remisión de Equipos")
EXCEL_FILE = "remision_equipos.xlsx"

class Equipo(BaseModel):
    categoria: str
    equipo: str
    cantidad: int

# Mapeo de columnas para el nuevo diseño de Excel (sin espacios en blanco)
CATEGORIAS_COLS = {
    "SONIDO": {"equipo": 1, "cantidad": 2},
    "VIDEO": {"equipo": 4, "cantidad": 5},
    "ILUMINACION Y ESTRUCTURA": {"equipo": 7, "cantidad": 8},
    "CABLEADO  ACCESORIOS Y OTROS": {"equipo": 10, "cantidad": 11},
    "PANTALLA LED": {"equipo": 13, "cantidad": 14}
}

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Remision"

        # Header customizado (Fila 2)
        ws.cell(row=2, column=1, value="Remisión de Salida de Equipos CD BOGOTA")
        
        ws.cell(row=4, column=1, value="Fecha de salida")
        ws.cell(row=4, column=3, value="Nombre del evento")
        ws.cell(row=4, column=8, value="Lugar del evento")

        # Categorias (Fila 6)
        ws.cell(row=6, column=1, value="SONIDO")
        ws.cell(row=6, column=4, value="VIDEO")
        ws.cell(row=6, column=7, value="ILUMINACION Y ESTRUCTURA")
        ws.cell(row=6, column=10, value="CABLEADO  ACCESORIOS Y OTROS")
        ws.cell(row=6, column=13, value="PANTALLA LED")

        # Subheaders (Fila 7)
        for col in [1, 4, 7, 10, 13]:
            ws.cell(row=7, column=col, value="EQUIPO")
            ws.cell(row=7, column=col+1, value="CANTIDAD")

        wb.save(EXCEL_FILE)

@app.on_event("startup")
def startup_event():
    init_excel()

@app.get("/", response_class=HTMLResponse)
def read_index():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()

@app.get("/equipos")
def obtener_equipos():
    if not os.path.exists(EXCEL_FILE):
        raise HTTPException(status_code=404, detail="Archivo Excel no encontrado")
    
    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active
    
    data = {}
    for cat_name, cols in CATEGORIAS_COLS.items():
        col_eq = cols["equipo"]
        col_cant = cols["cantidad"]
        equipos_cat = []
        
        # En el nuevo formato, los datos empiezan en la fila 8
        row = 8
        while True:
            equipo = ws.cell(row=row, column=col_eq).value
            cantidad = ws.cell(row=row, column=col_cant).value
            
            if equipo is None and cantidad is None:
                if ws.cell(row=row+1, column=col_eq).value is None:
                    break
            
            if equipo and isinstance(equipo, str) and ("Responsable" in equipo or "___" in equipo):
                break
                
            if equipo is not None:
                equipos_cat.append({"equipo": equipo, "cantidad": cantidad})
            row += 1
            
        data[cat_name] = equipos_cat
        
    return data

@app.post("/equipos/batch")
def agregar_equipos_batch(equipos_data: List[Equipo]):
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    
    # Escribir la fecha de hoy si es el primer dato que se agrega
    if ws.cell(row=4, column=2).value is None or str(ws.cell(row=4, column=2).value).strip() == "":
        ws.cell(row=4, column=2, value=datetime.now().strftime("%d/%m/%Y"))

    for equipo_data in equipos_data:
        categoria = equipo_data.categoria.upper()
        if categoria not in CATEGORIAS_COLS:
            continue
            
        cols = CATEGORIAS_COLS[categoria]
        col_eq = cols["equipo"]
        col_cant = cols["cantidad"]
        
        # Empezar a buscar desde la fila 8
        row = 8
        while True:
            val = ws.cell(row=row, column=col_eq).value
            if val is None or (isinstance(val, str) and ("Responsable" in val or "___" in val)):
                break
            row += 1
            
        equipo_formateado = equipo_data.equipo
        if equipo_formateado and isinstance(equipo_formateado, str) and len(equipo_formateado) > 0:
            equipo_formateado = equipo_formateado[0].upper() + equipo_formateado[1:]
        
        cell_eq = ws.cell(row=row, column=col_eq, value=equipo_formateado)
        cell_cant = ws.cell(row=row, column=col_cant, value=equipo_data.cantidad)
        
        # Ajustar texto y centrado
        cell_eq.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        cell_cant.alignment = Alignment(vertical='center', horizontal='center')
        
    try:
        wb.save(EXCEL_FILE)
    except PermissionError:
        raise HTTPException(status_code=400, detail="El archivo Excel está abierto. Ciérralo antes de intentar guardar.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
        
    return {"message": f"{len(equipos_data)} equipos agregados correctamente"}


@app.get("/descargar-excel")
def descargar_excel():
    if not os.path.exists(EXCEL_FILE):
        raise HTTPException(status_code=404, detail="El archivo Excel no existe o no ha sido creado an.")
    return FileResponse(path=EXCEL_FILE, filename="remision_equipos.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/abrir-excel")
def abrir_excel():
    try:
        if os.name == 'nt':  # Si es Windows
            os.startfile(EXCEL_FILE)
        return {"message": "Excel abierto correctamente"}
    except Exception as e:
        raise HTTPException(status_code=500, detail="No se pudo abrir el archivo Excel. Asegúrate de tener Excel instalado.")


@app.post("/borrar-listado")
def borrar_listado():
    try:
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        
        # Limpiar la fecha automática en la celda B4
        ws.cell(row=4, column=2).value = None

        # Limpiar datos desde la fila 8 hacia abajo en todas las columnas (incluyendo PANTALLA LED en M y N)
        max_row = max(ws.max_row, 35)
        max_col = max(ws.max_column, 14)
        for row in range(8, max_row + 1):
            for col in range(1, max_col + 1):
                ws.cell(row=row, column=col).value = None
                    
        wb.save(EXCEL_FILE)
        return {"status": "success", "message": "Listado borrado exitosamente"}
    except PermissionError:
        raise HTTPException(status_code=400, detail="El archivo Excel está abierto. Ciérralo antes de borrar el listado.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


class EmailRequest(BaseModel):
    remitente: str
    password: str
    destinatario: str
    asunto: str

@app.post("/enviar-correo")
def enviar_correo(req: EmailRequest):
    try:
        msg = EmailMessage()
        msg['Subject'] = req.asunto
        msg['From'] = req.remitente
        msg['To'] = req.destinatario
        msg.set_content("Adjunto encontrarás el Listado de Equipos de Eventos Corporativos.")

        # Verificar que el archivo existe
        if not os.path.exists(EXCEL_FILE):
            raise Exception("El archivo Excel aún no ha sido creado. Agrega al menos un equipo.")

        ctype, encoding = mimetypes.guess_type(EXCEL_FILE)
        if ctype is None or encoding is not None:
            ctype = 'application/octet-stream'
        maintype, subtype = ctype.split('/', 1)

        with open(EXCEL_FILE, 'rb') as fp:
            msg.add_attachment(fp.read(),
                               maintype=maintype,
                               subtype=subtype,
                               filename="Listado_Equipos.xlsx")

        # Conectar a Gmail
        with smtplib.SMTP_SSL('smtp.gmail.com', 465) as smtp:
            smtp.login(req.remitente, req.password)
            smtp.send_message(msg)

        return {"status": "success", "message": "Correo enviado"}
    except smtplib.SMTPAuthenticationError:
        raise HTTPException(status_code=401, detail="Error de autenticación. Verifica tu correo y asegúrate de estar usando una Contraseña de Aplicación.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
